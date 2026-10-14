import tkinter as tk
from tkinter import ttk, messagebox
import sqlite3
import os
import sys
import subprocess
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.units import mm

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side


# =========================================================
# PATHS
# =========================================================

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(PROJECT_DIR, "staff.db")
REPORT_DIR = os.path.join(PROJECT_DIR, "Leave_Reports")
SIGNATURE_DIR = os.path.join(PROJECT_DIR, "signatures")


# =========================================================
# COLORS
# =========================================================

BG = "#f4f7ff"
HEADER = "#071b49"
SIDEBAR = "#081633"
BLUE = "#2563eb"
PURPLE = "#7c3aed"
GREEN = "#16a34a"
RED = "#dc2626"
GRAY = "#64748b"
TEXT = "#172554"
WHITE = "#ffffff"
LINE = "#cbd5e1"


# =========================================================
# LEAVE LIMITS
# =========================================================

CL_LIMIT = 12
SL_LIMIT = 10


# =========================================================
# MONTHS
# =========================================================

MONTHS = [
    ("01", "January"),
    ("02", "February"),
    ("03", "March"),
    ("04", "April"),
    ("05", "May"),
    ("06", "June"),
    ("07", "July"),
    ("08", "August"),
    ("09", "September"),
    ("10", "October"),
    ("11", "November"),
    ("12", "December")
]

MONTH_NAMES = dict(MONTHS)


# =========================================================
# DATABASE
# =========================================================

def get_connection():
    return sqlite3.connect(DB_PATH)


def setup_database():

    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS leave_application (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            from_date TEXT NOT NULL,
            to_date TEXT NOT NULL,
            reason TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Pending',
            applied_date TEXT NOT NULL
        )
    """)

    # Check leave_type column
    cur.execute("PRAGMA table_info(leave_application)")
    columns = [row[1] for row in cur.fetchall()]

    if "leave_type" not in columns:

        cur.execute("""
            ALTER TABLE leave_application
            ADD COLUMN leave_type TEXT DEFAULT 'CL'
        """)

    conn.commit()
    conn.close()


# =========================================================
# MAIN CLASS
# =========================================================

class LeaveApplication:

    def __init__(self, root):

        self.root = root

        self.root.title(
            "Staff Management System - Leave Application"
        )

        self.root.geometry("1250x800")
        self.root.configure(bg=BG)

        setup_database()

        os.makedirs(
            REPORT_DIR,
            exist_ok=True
        )

        os.makedirs(
            SIGNATURE_DIR,
            exist_ok=True
        )

        self.selected_application_id = None

        self.build_ui()
        self.load_staff()
        self.load_applications()


    # =====================================================
    # BUILD UI
    # =====================================================

    def build_ui(self):

        # -------------------------------------------------
        # HEADER
        # -------------------------------------------------

        header = tk.Frame(
            self.root,
            bg=HEADER,
            height=70
        )

        header.pack(
            fill="x"
        )

        header.pack_propagate(False)

        tk.Label(
            header,
            text="STAFF MANAGEMENT SYSTEM",
            bg=HEADER,
            fg=WHITE,
            font=("Segoe UI", 20, "bold")
        ).pack(
            side="left",
            padx=25,
            pady=10
        )

        tk.Label(
            header,
            text="Leave Application",
            bg=HEADER,
            fg="#dbeafe",
            font=("Segoe UI", 11)
        ).pack(
            side="right",
            padx=25
        )


        # -------------------------------------------------
        # MAIN
        # -------------------------------------------------

        main = tk.Frame(
            self.root,
            bg=BG
        )

        main.pack(
            fill="both",
            expand=True,
            padx=18,
            pady=15
        )


        # =================================================
        # FORM
        # =================================================

        form = tk.LabelFrame(
            main,
            text=" Leave Application ",
            bg=WHITE,
            fg=HEADER,
            font=("Segoe UI", 12, "bold"),
            bd=1,
            relief="solid"
        )

        form.pack(
            fill="x",
            pady=(0, 12)
        )


        # -------------------------------------------------
        # STAFF
        # -------------------------------------------------

        tk.Label(
            form,
            text="Staff",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=0,
            column=0,
            padx=12,
            pady=10,
            sticky="w"
        )

        self.staff_var = tk.StringVar()

        self.staff_combo = ttk.Combobox(
            form,
            textvariable=self.staff_var,
            state="readonly",
            width=30
        )

        self.staff_combo.grid(
            row=0,
            column=1,
            padx=8,
            pady=10
        )

        self.staff_combo.bind(
            "<<ComboboxSelected>>",
            lambda e: self.load_applications()
        )


        # -------------------------------------------------
        # YEAR
        # -------------------------------------------------

        tk.Label(
            form,
            text="Year",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=0,
            column=2,
            padx=12,
            pady=10,
            sticky="w"
        )

        current_year = datetime.now().year

        self.year_var = tk.StringVar(
            value=str(current_year)
        )

        # Current year + next 5 years
        self.year_combo = ttk.Combobox(
            form,
            textvariable=self.year_var,
            state="readonly",
            width=10,
            values=[
                str(y)
                for y in range(
                    current_year,
                    current_year + 6
                )
            ]
        )

        self.year_combo.grid(
            row=0,
            column=3,
            padx=8,
            pady=10
        )

        self.year_combo.bind(
            "<<ComboboxSelected>>",
            lambda e: self.load_applications()
        )


        # -------------------------------------------------
        # MONTH
        # -------------------------------------------------

        tk.Label(
            form,
            text="Month",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=0,
            column=4,
            padx=12,
            pady=10,
            sticky="w"
        )

        current_month = datetime.now().strftime("%m")

        self.month_var = tk.StringVar(
            value=f"{current_month} - {MONTH_NAMES[current_month]}"
        )

        self.month_combo = ttk.Combobox(
            form,
            textvariable=self.month_var,
            state="readonly",
            width=18,
            values=[
                f"{num} - {name}"
                for num, name in MONTHS
            ]
        )

        self.month_combo.grid(
            row=0,
            column=5,
            padx=8,
            pady=10
        )

        self.month_combo.bind(
            "<<ComboboxSelected>>",
            lambda e: self.load_applications()
        )


        # -------------------------------------------------
        # FROM DATE
        # -------------------------------------------------

        tk.Label(
            form,
            text="From Date",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=1,
            column=0,
            padx=12,
            pady=10,
            sticky="w"
        )

        self.from_date_var = tk.StringVar()

        tk.Entry(
            form,
            textvariable=self.from_date_var,
            width=32,
            font=("Segoe UI", 10)
        ).grid(
            row=1,
            column=1,
            padx=8,
            pady=10
        )


        # -------------------------------------------------
        # TO DATE
        # -------------------------------------------------

        tk.Label(
            form,
            text="To Date",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=1,
            column=2,
            padx=12,
            pady=10,
            sticky="w"
        )

        self.to_date_var = tk.StringVar()

        tk.Entry(
            form,
            textvariable=self.to_date_var,
            width=20,
            font=("Segoe UI", 10)
        ).grid(
            row=1,
            column=3,
            padx=8,
            pady=10
        )


        # -------------------------------------------------
        # LEAVE TYPE
        # -------------------------------------------------

        tk.Label(
            form,
            text="Leave Type",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=1,
            column=4,
            padx=12,
            pady=10,
            sticky="w"
        )

        self.leave_type_var = tk.StringVar(
            value="CL"
        )

        self.leave_type_combo = ttk.Combobox(
            form,
            textvariable=self.leave_type_var,
            state="readonly",
            width=18,
            values=[
                "CL",
                "SL"
            ]
        )

        self.leave_type_combo.grid(
            row=1,
            column=5,
            padx=8,
            pady=10
        )


        # -------------------------------------------------
        # REASON
        # -------------------------------------------------

        tk.Label(
            form,
            text="Reason",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 10, "bold")
        ).grid(
            row=2,
            column=0,
            padx=12,
            pady=10,
            sticky="nw"
        )

        self.reason_text = tk.Text(
            form,
            height=3,
            width=80,
            font=("Segoe UI", 10),
            bd=1,
            relief="solid"
        )

        self.reason_text.grid(
            row=2,
            column=1,
            columnspan=5,
            padx=8,
            pady=10,
            sticky="w"
        )


        # -------------------------------------------------
        # BUTTONS
        # -------------------------------------------------

        button_frame = tk.Frame(
            form,
            bg=WHITE
        )

        button_frame.grid(
            row=3,
            column=0,
            columnspan=6,
            pady=12
        )


        tk.Button(
            button_frame,
            text="SAVE APPLICATION",
            bg=BLUE,
            fg=WHITE,
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padx=18,
            pady=8,
            command=self.save_application
        ).pack(
            side="left",
            padx=5
        )


        tk.Button(
            button_frame,
            text="UPDATE",
            bg=PURPLE,
            fg=WHITE,
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padx=18,
            pady=8,
            command=self.update_application
        ).pack(
            side="left",
            padx=5
        )


        tk.Button(
            button_frame,
            text="DELETE",
            bg=RED,
            fg=WHITE,
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padx=18,
            pady=8,
            command=self.delete_application
        ).pack(
            side="left",
            padx=5
        )


        tk.Button(
            button_frame,
            text="CLEAR",
            bg=GRAY,
            fg=WHITE,
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padx=18,
            pady=8,
            command=self.clear_form
        ).pack(
            side="left",
            padx=5
        )


        tk.Button(
            button_frame,
            text="EXPORT EXCEL",
            bg=GREEN,
            fg=WHITE,
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padx=18,
            pady=8,
            command=self.export_excel
        ).pack(
            side="left",
            padx=5
        )


        tk.Button(
            button_frame,
            text="OPEN REPORTS",
            bg=HEADER,
            fg=WHITE,
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padx=18,
            pady=8,
            command=self.open_report_folder
        ).pack(
            side="left",
            padx=5
        )


        # =================================================
        # SUMMARY CARDS
        # =================================================

        cards = tk.Frame(
            main,
            bg=BG
        )

        cards.pack(
            fill="x",
            pady=(0, 12)
        )


        # YEAR APPLICATIONS CARD REMOVED
        # Only 3 cards remain.


        self.month_app_card = self.create_card(
            cards,
            "MONTH APPLICATIONS",
            BLUE
        )


        self.month_cl_card = self.create_card(
            cards,
            "MONTH CL DAYS",
            PURPLE
        )


        self.month_sl_card = self.create_card(
            cards,
            "MONTH SL DAYS",
            GREEN
        )


        # =================================================
        # TABLE
        # =================================================

        table_frame = tk.Frame(
            main,
            bg=WHITE,
            bd=1,
            relief="solid"
        )

        table_frame.pack(
            fill="both",
            expand=True
        )


        columns = (
            "ID",
            "STAFF ID",
            "STAFF NAME",
            "FROM DATE",
            "TO DATE",
            "LEAVE TYPE",
            "REASON",
            "STATUS",
            "APPLIED DATE"
        )


        self.tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings"
        )


        for col in columns:

            self.tree.heading(
                col,
                text=col
            )


        # -------------------------------------------------
        # COLUMN WIDTH
        # -------------------------------------------------

        self.tree.column(
            "ID",
            width=55,
            anchor="center"
        )

        self.tree.column(
            "STAFF ID",
            width=75,
            anchor="center"
        )

        self.tree.column(
            "STAFF NAME",
            width=150
        )

        self.tree.column(
            "FROM DATE",
            width=100,
            anchor="center"
        )

        self.tree.column(
            "TO DATE",
            width=100,
            anchor="center"
        )

        self.tree.column(
            "LEAVE TYPE",
            width=90,
            anchor="center"
        )

        self.tree.column(
            "REASON",
            width=250
        )

        self.tree.column(
            "STATUS",
            width=90,
            anchor="center"
        )

        self.tree.column(
            "APPLIED DATE",
            width=120,
            anchor="center"
        )


        # -------------------------------------------------
        # SCROLLBAR
        # -------------------------------------------------

        scrollbar_y = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.tree.yview
        )

        self.tree.configure(
            yscrollcommand=scrollbar_y.set
        )


        self.tree.pack(
            side="left",
            fill="both",
            expand=True
        )


        scrollbar_y.pack(
            side="right",
            fill="y"
        )


        self.tree.bind(
            "<Double-1>",
            self.select_record
        )


    # =====================================================
    # CREATE CARD
    # =====================================================

    def create_card(
        self,
        parent,
        title,
        color
    ):

        card = tk.Frame(
            parent,
            bg=WHITE,
            bd=1,
            relief="solid",
            height=105
        )

        card.pack(
            side="left",
            fill="both",
            expand=True,
            padx=5
        )

        card.pack_propagate(False)


        tk.Label(
            card,
            text=title,
            bg=WHITE,
            fg=GRAY,
            font=("Segoe UI", 10, "bold")
        ).pack(
            pady=(15, 5)
        )


        value_label = tk.Label(
            card,
            text="0",
            bg=WHITE,
            fg=color,
            font=("Segoe UI", 25, "bold")
        )

        value_label.pack()

        return value_label


    # =====================================================
    # LOAD STAFF
    # =====================================================

    def load_staff(self):

        conn = get_connection()
        cur = conn.cursor()

        try:

            cur.execute("""
                SELECT id, name
                FROM staff
                ORDER BY id
            """)

            rows = cur.fetchall()

        except sqlite3.Error:

            rows = []

        conn.close()


        staff_values = [
            f"{row[0]} - {row[1]}"
            for row in rows
        ]


        self.staff_combo["values"] = staff_values


        if staff_values:

            self.staff_combo.current(0)


    # =====================================================
    # GET STAFF ID
    # =====================================================

    def get_selected_staff_id(self):

        value = self.staff_var.get().strip()

        if not value:
            return None

        try:

            return int(
                value.split("-")[0].strip()
            )

        except:

            return None


    # =====================================================
    # GET STAFF NAME
    # =====================================================

    def get_staff_name(
        self,
        staff_id
    ):

        conn = get_connection()
        cur = conn.cursor()

        cur.execute(
            "SELECT name FROM staff WHERE id=?",
            (staff_id,)
        )

        row = cur.fetchone()

        conn.close()

        return row[0] if row else "Unknown"


    # =====================================================
    # GET MONTH NUMBER
    # =====================================================

    def get_month_number(self):

        value = self.month_var.get()

        if not value:
            return ""

        return value.split("-")[0].strip()


    # =====================================================
    # CALCULATE DAYS
    # =====================================================

    def calculate_days(
        self,
        from_date,
        to_date
    ):

        try:

            d1 = datetime.strptime(
                from_date,
                "%Y-%m-%d"
            )

            d2 = datetime.strptime(
                to_date,
                "%Y-%m-%d"
            )

            return (
                d2 - d1
            ).days + 1

        except:

            return 0


    # =====================================================
    # GET USED LEAVE
    # =====================================================

    def get_used_leave(
        self,
        staff_id,
        leave_type,
        year
    ):

        conn = get_connection()
        cur = conn.cursor()

        cur.execute("""
            SELECT from_date, to_date
            FROM leave_application
            WHERE staff_id=?
            AND leave_type=?
            AND substr(from_date,1,4)=?
        """, (
            staff_id,
            leave_type,
            str(year)
        ))

        rows = cur.fetchall()

        conn.close()


        total = 0

        for from_date, to_date in rows:

            total += self.calculate_days(
                from_date,
                to_date
            )


        return total


    # =====================================================
    # MONTH APPLICATION COUNT
    # =====================================================

    def get_month_application_count(
        self,
        staff_id,
        year,
        month
    ):

        conn = get_connection()
        cur = conn.cursor()

        cur.execute("""
            SELECT COUNT(*)
            FROM leave_application
            WHERE staff_id=?
            AND substr(from_date,1,4)=?
            AND substr(from_date,6,2)=?
        """, (
            staff_id,
            str(year),
            str(month)
        ))

        result = cur.fetchone()[0]

        conn.close()

        return result


    # =====================================================
    # MONTH LEAVE DAYS
    # =====================================================

    def get_month_leave_days(
        self,
        staff_id,
        leave_type,
        year,
        month
    ):

        conn = get_connection()
        cur = conn.cursor()

        cur.execute("""
            SELECT from_date, to_date
            FROM leave_application
            WHERE staff_id=?
            AND leave_type=?
            AND substr(from_date,1,4)=?
            AND substr(from_date,6,2)=?
        """, (
            staff_id,
            leave_type,
            str(year),
            str(month)
        ))

        rows = cur.fetchall()

        conn.close()


        total = 0

        for from_date, to_date in rows:

            total += self.calculate_days(
                from_date,
                to_date
            )


        return total


    # =====================================================
    # LOAD APPLICATIONS
    # =====================================================

    def load_applications(self):

        # Clear table
        for item in self.tree.get_children():

            self.tree.delete(item)


        staff_id = self.get_selected_staff_id()

        year = self.year_var.get().strip()

        month = self.get_month_number()


        # -------------------------------------------------
        # UPDATE SUMMARY CARDS
        # -------------------------------------------------

        if staff_id and year and month:

            month_applications = (
                self.get_month_application_count(
                    staff_id,
                    year,
                    month
                )
            )

            month_cl_days = (
                self.get_month_leave_days(
                    staff_id,
                    "CL",
                    year,
                    month
                )
            )

            month_sl_days = (
                self.get_month_leave_days(
                    staff_id,
                    "SL",
                    year,
                    month
                )
            )

        else:

            month_applications = 0
            month_cl_days = 0
            month_sl_days = 0


        self.month_app_card.config(
            text=str(month_applications)
        )

        self.month_cl_card.config(
            text=str(month_cl_days)
        )

        self.month_sl_card.config(
            text=str(month_sl_days)
        )


        # -------------------------------------------------
        # DATABASE QUERY
        # -------------------------------------------------

        conn = get_connection()
        cur = conn.cursor()


        query = """
            SELECT
                l.id,
                l.staff_id,
                s.name,
                l.from_date,
                l.to_date,
                l.leave_type,
                l.reason,
                l.status,
                l.applied_date
            FROM leave_application l
            LEFT JOIN staff s
            ON l.staff_id=s.id
            WHERE substr(l.from_date,1,4)=?
            AND substr(l.from_date,6,2)=?
        """


        params = [
            year,
            month
        ]


        if staff_id:

            query += """
                AND l.staff_id=?
            """

            params.append(
                staff_id
            )


        query += """
            ORDER BY l.id DESC
        """


        cur.execute(
            query,
            params
        )


        rows = cur.fetchall()

        conn.close()


        # -------------------------------------------------
        # INSERT TABLE DATA
        #
        # IMPORTANT:
        # Database ID remains hidden internally.
        # Display ID = 1,2,3,4...
        # -------------------------------------------------

        for serial_no, row in enumerate(
            rows,
            start=1
        ):

            actual_database_id = row[0]

            display_values = (
                serial_no,       # Display Serial ID
                row[1],          # Staff ID
                row[2],          # Staff Name
                row[3],          # From Date
                row[4],          # To Date
                row[5],          # Leave Type
                row[6],          # Reason
                row[7],          # Status
                row[8]           # Applied Date
            )


            # iid stores actual database ID
            # but is NOT displayed
            self.tree.insert(
                "",
                "end",
                iid=str(actual_database_id),
                values=display_values
            )


    # =====================================================
    # VALIDATE DATES
    # =====================================================

    def validate_dates(
        self,
        from_date,
        to_date
    ):

        try:

            d1 = datetime.strptime(
                from_date,
                "%Y-%m-%d"
            )

            d2 = datetime.strptime(
                to_date,
                "%Y-%m-%d"
            )


            if d2 < d1:

                messagebox.showerror(
                    "Invalid Date",
                    "To Date cannot be before From Date."
                )

                return False


            return True


        except:

            messagebox.showerror(
                "Invalid Date",
                "Please enter date in YYYY-MM-DD format."
            )

            return False


    # =====================================================
    # SAVE APPLICATION
    # =====================================================

    def save_application(self):

        staff_id = self.get_selected_staff_id()

        from_date = (
            self.from_date_var.get().strip()
        )

        to_date = (
            self.to_date_var.get().strip()
        )

        leave_type = (
            self.leave_type_var.get().strip()
        )

        reason = (
            self.reason_text.get(
                "1.0",
                "end"
            ).strip()
        )


        if not staff_id:

            messagebox.showwarning(
                "Required",
                "Please select staff."
            )

            return


        if not from_date or not to_date:

            messagebox.showwarning(
                "Required",
                "Please enter From Date and To Date."
            )

            return


        if not reason:

            messagebox.showwarning(
                "Required",
                "Please enter reason."
            )

            return


        if not self.validate_dates(
            from_date,
            to_date
        ):

            return


        year = from_date[:4]


        # -------------------------------------------------
        # LEAVE LIMIT CHECK
        # -------------------------------------------------

        used = self.get_used_leave(
            staff_id,
            leave_type,
            year
        )


        requested_days = self.calculate_days(
            from_date,
            to_date
        )


        limit = (
            CL_LIMIT
            if leave_type == "CL"
            else SL_LIMIT
        )


        if used + requested_days > limit:

            messagebox.showerror(
                "Leave Limit Exceeded",
                f"{leave_type} annual limit is {limit} days.\n\n"
                f"Already Used: {used} days\n"
                f"Requested: {requested_days} days\n"
                f"Remaining: {max(0, limit - used)} days"
            )

            return


        applied_date = datetime.now().strftime(
            "%Y-%m-%d"
        )


        conn = get_connection()
        cur = conn.cursor()


        cur.execute("""
            INSERT INTO leave_application
            (
                staff_id,
                from_date,
                to_date,
                reason,
                status,
                applied_date,
                leave_type
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            staff_id,
            from_date,
            to_date,
            reason,
            "Pending",
            applied_date,
            leave_type
        ))


        application_id = cur.lastrowid


        conn.commit()
        conn.close()


        messagebox.showinfo(
            "Success",
            "Leave Application Saved Successfully!"
        )


        # Set year/month according to application
        self.year_var.set(
            from_date[:4]
        )


        month_no = from_date[5:7]


        if month_no in MONTH_NAMES:

            self.month_var.set(
                f"{month_no} - {MONTH_NAMES[month_no]}"
            )


        # Generate PDF
        self.generate_report(
            application_id
        )


        self.load_applications()


        self.clear_form(
            keep_staff=True
        )


    # =====================================================
    # SELECT RECORD
    # =====================================================

    def select_record(
        self,
        event=None
    ):

        selected = self.tree.selection()


        if not selected:
            return


        # -------------------------------------------------
        # IMPORTANT:
        # Treeview IID contains actual database ID.
        # Display ID is only serial number.
        # -------------------------------------------------

        actual_database_id = int(
            selected[0]
        )


        self.selected_application_id = (
            actual_database_id
        )


        values = self.tree.item(
            selected[0],
            "values"
        )


        if not values:
            return


        staff_id = values[1]
        staff_name = values[2]


        self.staff_var.set(
            f"{staff_id} - {staff_name}"
        )


        self.from_date_var.set(
            values[3]
        )


        self.to_date_var.set(
            values[4]
        )


        self.leave_type_var.set(
            values[5]
        )


        self.reason_text.delete(
            "1.0",
            "end"
        )


        self.reason_text.insert(
            "1.0",
            values[6]
        )


    # =====================================================
    # UPDATE APPLICATION
    # =====================================================

    def update_application(self):

        if not self.selected_application_id:

            messagebox.showwarning(
                "Select Record",
                "Please select an application first."
            )

            return


        staff_id = self.get_selected_staff_id()

        from_date = (
            self.from_date_var.get().strip()
        )

        to_date = (
            self.to_date_var.get().strip()
        )

        leave_type = (
            self.leave_type_var.get().strip()
        )

        reason = (
            self.reason_text.get(
                "1.0",
                "end"
            ).strip()
        )


        if (
            not staff_id
            or not from_date
            or not to_date
            or not reason
        ):

            messagebox.showwarning(
                "Required",
                "Please fill all required fields."
            )

            return


        if not self.validate_dates(
            from_date,
            to_date
        ):

            return


        conn = get_connection()
        cur = conn.cursor()


        cur.execute("""
            UPDATE leave_application
            SET
                staff_id=?,
                from_date=?,
                to_date=?,
                reason=?,
                leave_type=?
            WHERE id=?
        """, (
            staff_id,
            from_date,
            to_date,
            reason,
            leave_type,
            self.selected_application_id
        ))


        conn.commit()
        conn.close()


        messagebox.showinfo(
            "Updated",
            "Leave Application Updated Successfully!"
        )


        self.generate_report(
            self.selected_application_id
        )


        self.load_applications()


        self.clear_form(
            keep_staff=True
        )


    # =====================================================
    # DELETE APPLICATION
    # =====================================================

    def delete_application(self):

        if not self.selected_application_id:

            messagebox.showwarning(
                "Select Record",
                "Please select an application first."
            )

            return


        confirm = messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete this application?"
        )


        if not confirm:
            return


        conn = get_connection()
        cur = conn.cursor()


        cur.execute("""
            DELETE FROM leave_application
            WHERE id=?
        """, (
            self.selected_application_id,
        ))


        conn.commit()
        conn.close()


        messagebox.showinfo(
            "Deleted",
            "Leave Application Deleted Successfully!"
        )


        self.selected_application_id = None


        self.load_applications()


        self.clear_form(
            keep_staff=True
        )


    # =====================================================
    # CLEAR FORM
    # =====================================================

    def clear_form(
        self,
        keep_staff=False
    ):

        if not keep_staff:

            self.staff_var.set("")


        self.from_date_var.set("")
        self.to_date_var.set("")


        self.leave_type_var.set(
            "CL"
        )


        self.reason_text.delete(
            "1.0",
            "end"
        )


        self.selected_application_id = None


    # =====================================================
    # GENERATE PDF
    # =====================================================

    def generate_report(
        self,
        application_id
    ):

        conn = get_connection()
        cur = conn.cursor()


        cur.execute("""
            SELECT
                l.id,
                l.staff_id,
                s.name,
                l.from_date,
                l.to_date,
                l.leave_type,
                l.reason,
                l.status,
                l.applied_date
            FROM leave_application l
            LEFT JOIN staff s
            ON l.staff_id=s.id
            WHERE l.id=?
        """, (
            application_id,
        ))


        row = cur.fetchone()


        conn.close()


        if not row:
            return


        (
            app_id,
            staff_id,
            staff_name,
            from_date,
            to_date,
            leave_type,
            reason,
            status,
            applied_date
        ) = row


        safe_name = "".join(
            c
            for c in staff_name
            if c.isalnum()
            or c in (
                " ",
                "_",
                "-"
            )
        ).strip()


        staff_report_dir = os.path.join(
            REPORT_DIR,
            safe_name
        )


        os.makedirs(
            staff_report_dir,
            exist_ok=True
        )


        pdf_path = os.path.join(
            staff_report_dir,
            f"Leave_Application_{app_id}.pdf"
        )


        c = canvas.Canvas(
            pdf_path,
            pagesize=A4
        )


        width, height = A4


        # -------------------------------------------------
        # HEADER
        # -------------------------------------------------

        c.setFillColor(
            colors.HexColor(HEADER)
        )


        c.rect(
            0,
            height - 35 * mm,
            width,
            35 * mm,
            fill=1,
            stroke=0
        )


        c.setFillColor(
            colors.white
        )


        c.setFont(
            "Helvetica-Bold",
            18
        )


        c.drawCentredString(
            width / 2,
            height - 14 * mm,
            "STAFF MANAGEMENT SYSTEM"
        )


        c.setFont(
            "Helvetica",
            9
        )


        c.drawCentredString(
            width / 2,
            height - 21 * mm,
            "Smart • Simple • Efficient Staff Management"
        )


        # -------------------------------------------------
        # TITLE
        # -------------------------------------------------

        c.setFillColor(
            colors.HexColor(TEXT)
        )


        c.setFont(
            "Helvetica-Bold",
            16
        )


        c.drawCentredString(
            width / 2,
            height - 52 * mm,
            "LEAVE APPLICATION"
        )


        # -------------------------------------------------
        # DETAILS
        # -------------------------------------------------

        y = height - 70 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "Application ID:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            str(app_id)
        )


        y -= 10 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "Staff Name:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            staff_name
        )


        y -= 10 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "Leave Type:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            leave_type
        )


        y -= 10 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "From Date:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            from_date
        )


        y -= 10 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "To Date:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            to_date
        )


        y -= 10 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "Applied Date:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            applied_date
        )


        y -= 10 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "Status:"
        )


        c.setFont(
            "Helvetica",
            10
        )


        c.drawString(
            65 * mm,
            y,
            status
        )


        # -------------------------------------------------
        # REASON
        # -------------------------------------------------

        y -= 18 * mm


        c.setFont(
            "Helvetica-Bold",
            10
        )


        c.drawString(
            25 * mm,
            y,
            "Reason:"
        )


        y -= 8 * mm


        c.setFont(
            "Helvetica",
            10
        )


        words = reason.split()

        line = ""


        for word in words:

            test_line = (
                line + " " + word
            ).strip()


            if c.stringWidth(
                test_line,
                "Helvetica",
                10
            ) > 150 * mm:

                c.drawString(
                    30 * mm,
                    y,
                    line
                )


                y -= 6 * mm


                line = word

            else:

                line = test_line


        if line:

            c.drawString(
                30 * mm,
                y,
                line
            )


        # -------------------------------------------------
        # PRINCIPAL SIGNATURE ONLY
        # -------------------------------------------------

        signature_path = os.path.join(
            SIGNATURE_DIR,
            "khardekar_sir.png"
        )


        if os.path.exists(
            signature_path
        ):

            try:

                c.drawImage(
                    signature_path,
                    width - 75 * mm,
                    35 * mm,
                    width=45 * mm,
                    height=18 * mm,
                    preserveAspectRatio=True,
                    mask="auto"
                )

            except Exception:

                pass


        c.setFont(
            "Helvetica-Bold",
            9
        )


        c.drawCentredString(
            width - 52.5 * mm,
            30 * mm,
            "Principal / Authorized Signature"
        )


        # -------------------------------------------------
        # FOOTER
        # -------------------------------------------------

        c.setStrokeColor(
            colors.HexColor(LINE)
        )


        c.line(
            20 * mm,
            20 * mm,
            width - 20 * mm,
            20 * mm
        )


        c.setFillColor(
            colors.HexColor(GRAY)
        )


        c.setFont(
            "Helvetica",
            8
        )


        c.drawCentredString(
            width / 2,
            14 * mm,
            "Official Leave Application"
        )


        c.save()


        return pdf_path


    # =====================================================
    # EXPORT EXCEL
    # =====================================================

    def export_excel(self):

        staff_id = self.get_selected_staff_id()


        if not staff_id:

            messagebox.showwarning(
                "Select Staff",
                "Please select staff first."
            )

            return


        year = self.year_var.get()

        staff_name = self.get_staff_name(
            staff_id
        )


        safe_name = "".join(
            c
            for c in staff_name
            if c.isalnum()
            or c in (
                " ",
                "_",
                "-"
            )
        ).strip()


        staff_report_dir = os.path.join(
            REPORT_DIR,
            safe_name
        )


        os.makedirs(
            staff_report_dir,
            exist_ok=True
        )


        excel_path = os.path.join(
            staff_report_dir,
            f"Leave_Report_{safe_name}_{year}.xlsx"
        )


        conn = get_connection()
        cur = conn.cursor()


        cur.execute("""
            SELECT
                l.id,
                l.staff_id,
                s.name,
                l.from_date,
                l.to_date,
                l.leave_type,
                l.reason,
                l.status,
                l.applied_date
            FROM leave_application l
            LEFT JOIN staff s
            ON l.staff_id=s.id
            WHERE l.staff_id=?
            AND substr(l.from_date,1,4)=?
            ORDER BY l.id
        """, (
            staff_id,
            year
        ))


        rows = cur.fetchall()


        conn.close()


        wb = Workbook()

        ws = wb.active

        ws.title = "Leave Report"


        headers = [
            "ID",
            "Staff ID",
            "Staff Name",
            "From Date",
            "To Date",
            "Leave Type",
            "Reason",
            "Status",
            "Applied Date"
        ]


        ws.append(
            headers
        )


        for serial_no, row in enumerate(
            rows,
            start=1
        ):

            ws.append([
                serial_no,
                row[1],
                row[2],
                row[3],
                row[4],
                row[5],
                row[6],
                row[7],
                row[8]
            ])


        # -------------------------------------------------
        # EXCEL STYLE
        # -------------------------------------------------

        header_font = Font(
            bold=True,
            color="FFFFFF"
        )


        for cell in ws[1]:

            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


        thin = Side(
            style="thin",
            color="CCCCCC"
        )


        for row in ws.iter_rows():

            for cell in row:

                cell.border = Border(
                    left=thin,
                    right=thin,
                    top=thin,
                    bottom=thin
                )

                cell.alignment = Alignment(
                    vertical="center"
                )


        widths = {
            "A": 8,
            "B": 10,
            "C": 25,
            "D": 15,
            "E": 15,
            "F": 14,
            "G": 40,
            "H": 15,
            "I": 18
        }


        for col, width in widths.items():

            ws.column_dimensions[
                col
            ].width = width


        wb.save(
            excel_path
        )


        messagebox.showinfo(
            "Excel Exported",
            f"Excel report created successfully.\n\n{excel_path}"
        )


    # =====================================================
    # OPEN REPORT FOLDER
    # =====================================================

    def open_report_folder(self):

        os.makedirs(
            REPORT_DIR,
            exist_ok=True
        )


        try:

            if sys.platform.startswith("win"):

                os.startfile(
                    REPORT_DIR
                )

            elif sys.platform == "darwin":

                subprocess.Popen(
                    ["open", REPORT_DIR]
                )

            else:

                subprocess.Popen(
                    ["xdg-open", REPORT_DIR]
                )


        except Exception as e:

            messagebox.showerror(
                "Error",
                str(e)
            )


# =========================================================
# RUN
# =========================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = LeaveApplication(
        root
    )

    root.mainloop()