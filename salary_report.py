import tkinter as tk
from tkinter import ttk, messagebox
import sqlite3
import os
import calendar
import subprocess
import sys
from datetime import datetime

# ============================================================
# STAFF MANAGEMENT SYSTEM - SALARY REPORT
# ============================================================
# This file works with:
#   staff table      -> id, name, salary
#   attendance table -> id, staff_id, date, status
#
# Attendance statuses:
#   Present, Absent, CL, SL
#
# Salary calculation:
#   Daily Rate = Monthly Salary / Calendar Days of that month
#   Absent Cut = Absent Days * Daily Rate
#   Payable = Monthly Salary - Absent Cut
#   Present / CL / SL are treated as paid.
#
# This program also creates:
#   salary_records
# table automatically for saving month-wise salary records.
# ============================================================


# ============================================================
# PATHS
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "staff.db")
REPORT_DIR = os.path.join(BASE_DIR, "Salary_Reports")

os.makedirs(REPORT_DIR, exist_ok=True)


# ============================================================
# COLORS
# ============================================================

BG = "#f4f7ff"
HEADER = "#071b49"
BLUE = "#2563eb"
GREEN = "#16a34a"
RED = "#dc2626"
PURPLE = "#7c3aed"
ORANGE = "#f59e0b"
WHITE = "#ffffff"
TEXT = "#172554"
GRAY = "#64748b"
LINE = "#cbd5e1"
LIGHT_BLUE = "#dbeafe"
LIGHT_GREEN = "#dcfce7"
LIGHT_RED = "#fee2e2"


# ============================================================
# DATABASE
# ============================================================

def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def get_staff_list():
    """
    Read staff from existing staff table.
    Expected columns: id, name, salary.
    A fallback for staff_name is included because some projects
    use staff_name instead of name.
    """
    conn = get_connection()
    cur = conn.cursor()

    try:
        cur.execute(
            "SELECT id, name, salary FROM staff ORDER BY id"
        )
        rows = cur.fetchall()
    except sqlite3.OperationalError:
        cur.execute(
            "SELECT id, staff_name, salary FROM staff ORDER BY id"
        )
        rows = cur.fetchall()
    finally:
        conn.close()

    return rows


def setup_salary_records_table():
    """
    Creates a permanent month-wise salary record table.
    One staff member can have one record for each month/year.
    """
    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS salary_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            year INTEGER NOT NULL,
            month INTEGER NOT NULL,
            monthly_salary REAL NOT NULL DEFAULT 0,
            calendar_days INTEGER NOT NULL DEFAULT 0,
            present_days INTEGER NOT NULL DEFAULT 0,
            absent_days INTEGER NOT NULL DEFAULT 0,
            cl_days INTEGER NOT NULL DEFAULT 0,
            sl_days INTEGER NOT NULL DEFAULT 0,
            absent_cut REAL NOT NULL DEFAULT 0,
            payable_salary REAL NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            UNIQUE(staff_id, year, month)
        )
        """
    )

    conn.commit()
    conn.close()


def get_attendance(staff_id, year, month=None):
    """
    Get attendance for one staff member.

    If month is supplied, only that month is returned.
    """
    conn = get_connection()
    cur = conn.cursor()

    if month is None:
        cur.execute(
            """
            SELECT date, status
            FROM attendance
            WHERE staff_id = ?
              AND date LIKE ?
            ORDER BY date
            """,
            (staff_id, f"{year}-%")
        )
    else:
        prefix = f"{year}-{month:02d}-%"
        cur.execute(
            """
            SELECT date, status
            FROM attendance
            WHERE staff_id = ?
              AND date LIKE ?
            ORDER BY date
            """,
            (staff_id, prefix)
        )

    rows = cur.fetchall()
    conn.close()
    return rows


# ============================================================
# SALARY CALCULATION
# ============================================================

def calculate_month(records, monthly_salary, year, month):
    """
    Calculate one month's salary.

    Example:
        Monthly salary = 15000
        September = 30 days
        Absent = 2

        Daily rate = 15000 / 30 = 500
        Absent cut = 2 * 500 = 1000
        Payable = 15000 - 1000 = 14000
    """

    days_in_month = calendar.monthrange(year, month)[1]

    daily_rate = (
        monthly_salary / days_in_month
        if days_in_month > 0
        else 0
    )

    present = 0
    absent = 0
    cl = 0
    sl = 0

    for date_value, status in records:
        try:
            dt = datetime.strptime(
                str(date_value).strip(),
                "%Y-%m-%d"
            )
        except ValueError:
            continue

        if dt.year != year or dt.month != month:
            continue

        status = str(status).strip().lower()

        if status == "present":
            present += 1
        elif status == "absent":
            absent += 1
        elif status == "cl":
            cl += 1
        elif status == "sl":
            sl += 1

    absent_cut = absent * daily_rate
    payable = max(monthly_salary - absent_cut, 0)

    return {
        "month": calendar.month_name[month],
        "month_no": month,
        "days": days_in_month,
        "present": present,
        "absent": absent,
        "cl": cl,
        "sl": sl,
        "recorded": present + absent + cl + sl,
        "daily_rate": daily_rate,
        "salary": monthly_salary,
        "deduction": absent_cut,
        "payable": payable,
    }


def calculate_year(records, monthly_salary, year):
    monthly = []

    for month in range(1, 13):
        monthly.append(
            calculate_month(
                records,
                monthly_salary,
                year,
                month
            )
        )

    return {
        "monthly": monthly,
        "year_salary": monthly_salary * 12,
        "present": sum(x["present"] for x in monthly),
        "absent": sum(x["absent"] for x in monthly),
        "cl": sum(x["cl"] for x in monthly),
        "sl": sum(x["sl"] for x in monthly),
        "recorded": sum(x["recorded"] for x in monthly),
        "total_deduction": sum(
            x["deduction"] for x in monthly
        ),
        "total_payable": sum(
            x["payable"] for x in monthly
        ),
    }


# ============================================================
# DATABASE SAVE / INSERT
# ============================================================

def save_month_record(staff_id, year, item):
    """
    Insert a month-wise salary record.
    If the same staff/year/month already exists, update it.
    """
    conn = get_connection()
    cur = conn.cursor()

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cur.execute(
        """
        INSERT INTO salary_records (
            staff_id,
            year,
            month,
            monthly_salary,
            calendar_days,
            present_days,
            absent_days,
            cl_days,
            sl_days,
            absent_cut,
            payable_salary,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(staff_id, year, month)
        DO UPDATE SET
            monthly_salary = excluded.monthly_salary,
            calendar_days = excluded.calendar_days,
            present_days = excluded.present_days,
            absent_days = excluded.absent_days,
            cl_days = excluded.cl_days,
            sl_days = excluded.sl_days,
            absent_cut = excluded.absent_cut,
            payable_salary = excluded.payable_salary,
            updated_at = excluded.updated_at
        """,
        (
            staff_id,
            year,
            item["month_no"],
            item["salary"],
            item["days"],
            item["present"],
            item["absent"],
            item["cl"],
            item["sl"],
            item["deduction"],
            item["payable"],
            now,
            now,
        )
    )

    conn.commit()
    conn.close()


def save_year_records(staff_id, year, summary):
    """
    Save all 12 months for selected staff.
    """
    for item in summary["monthly"]:
        save_month_record(
            staff_id,
            year,
            item
        )


def get_saved_records(staff_id, year):
    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        SELECT
            month,
            monthly_salary,
            calendar_days,
            present_days,
            absent_days,
            cl_days,
            sl_days,
            absent_cut,
            payable_salary
        FROM salary_records
        WHERE staff_id = ?
          AND year = ?
        ORDER BY month
        """,
        (staff_id, year)
    )

    rows = cur.fetchall()
    conn.close()

    return rows


# ============================================================
# PDF
# ============================================================

def create_pdf(
    staff_name,
    staff_id,
    year,
    monthly_salary,
    summary,
    pdf_path
):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import (
            getSampleStyleSheet,
            ParagraphStyle
        )
        from reportlab.lib.enums import TA_CENTER
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle
        )
    except ImportError:
        raise RuntimeError(
            "reportlab is not installed.\n\n"
            "Run:\n"
            "py -m pip install reportlab"
        )

    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()

    title = ParagraphStyle(
        "SalaryTitle",
        parent=styles["Title"],
        fontSize=19,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#071b49"),
        spaceAfter=5
    )

    center = ParagraphStyle(
        "SalaryCenter",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=9
    )

    story = [
        Paragraph(
            "STAFF MANAGEMENT SYSTEM",
            title
        ),
        Paragraph(
            "SALARY & ATTENDANCE PAYMENT REPORT",
            center
        ),
        Spacer(1, 12),
        Paragraph(
            f"<b>Staff:</b> {staff_name} &nbsp;&nbsp;&nbsp; "
            f"<b>Staff ID:</b> {staff_id} &nbsp;&nbsp;&nbsp; "
            f"<b>Year:</b> {year} &nbsp;&nbsp;&nbsp; "
            f"<b>Monthly Salary:</b> "
            f"Rs. {monthly_salary:,.2f}",
            styles["Normal"]
        ),
        Spacer(1, 12)
    ]

    headers = [
        "Month",
        "Days",
        "Recorded",
        "Present",
        "Absent",
        "CL",
        "SL",
        "Monthly Salary",
        "Absent Cut",
        "Payable"
    ]

    data = [headers]

    for item in summary["monthly"]:
        data.append([
            item["month"],
            item["days"],
            item["recorded"],
            item["present"],
            item["absent"],
            item["cl"],
            item["sl"],
            f"Rs. {item['salary']:,.2f}",
            f"Rs. {item['deduction']:,.2f}",
            f"Rs. {item['payable']:,.2f}"
        ])

    data.append([
        "YEAR TOTAL",
        "",
        summary["recorded"],
        summary["present"],
        summary["absent"],
        summary["cl"],
        summary["sl"],
        f"Rs. {summary['year_salary']:,.2f}",
        f"Rs. {summary['total_deduction']:,.2f}",
        f"Rs. {summary['total_payable']:,.2f}"
    ])

    table = Table(
        data,
        colWidths=[
            63, 35, 45, 45, 45,
            35, 35, 80, 75, 80
        ],
        repeatRows=1
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#071b49")
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor("#cbd5e1")
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -2),
                [
                    colors.white,
                    colors.HexColor("#f8fafc")
                ]
            ),
            (
                "BACKGROUND",
                (0, -1),
                (-1, -1),
                colors.HexColor("#dbeafe")
            ),
            (
                "FONTNAME",
                (0, -1),
                (-1, -1),
                "Helvetica-Bold"
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5
            ),
        ])
    )

    story.append(table)
    story.append(Spacer(1, 14))

    story.append(
        Paragraph(
            f"<b>YEAR PAYMENT SUMMARY</b><br/>"
            f"Year Salary = Rs. "
            f"{summary['year_salary']:,.2f}"
            f" &nbsp; | &nbsp; "
            f"Absent Cut = Rs. "
            f"{summary['total_deduction']:,.2f}"
            f" &nbsp; | &nbsp; "
            f"Final Year Payment = Rs. "
            f"{summary['total_payable']:,.2f}",
            styles["Normal"]
        )
    )

    story.append(Spacer(1, 8))

    story.append(
        Paragraph(
            "Calculation: "
            "Absent Cut = Absent Days × "
            "(Monthly Salary ÷ Calendar Days). "
            "Present, CL and SL are paid.",
            styles["Normal"]
        )
    )

    story.append(Spacer(1, 6))

    story.append(
        Paragraph(
            "Note: Salary is calculated from attendance records "
            "stored in the attendance table.",
            styles["Normal"]
        )
    )

    doc.build(story)


# ============================================================
# EXCEL
# ============================================================

def create_excel(
    staff_name,
    staff_id,
    year,
    monthly_salary,
    summary,
    excel_path
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Font,
            PatternFill,
            Border,
            Side,
            Alignment
        )
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import (
            Table,
            TableStyleInfo
        )
    except ImportError:
        raise RuntimeError(
            "openpyxl is not installed.\n\n"
            "Run:\n"
            "py -m pip install openpyxl"
        )

    wb = Workbook()

    ws = wb.active
    ws.title = "Salary Report"

    title_fill = PatternFill(
        "solid",
        fgColor="071B49"
    )

    header_fill = PatternFill(
        "solid",
        fgColor="2563EB"
    )

    total_fill = PatternFill(
        "solid",
        fgColor="DBEAFE"
    )

    thin = Side(
        style="thin",
        color="CBD5E1"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    # Title
    ws.merge_cells("A1:J1")
    ws["A1"] = (
        "STAFF MANAGEMENT SYSTEM - SALARY REPORT"
    )
    ws["A1"].font = Font(
        size=18,
        bold=True,
        color="FFFFFF"
    )
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    # Staff information
    info = [
        ("A2", "Staff", "B2", staff_name),
        ("D2", "Staff ID", "E2", staff_id),
        ("G2", "Year", "H2", year),
        ("A3", "Monthly Salary", "B3", monthly_salary),
        (
            "D3",
            "Year Salary",
            "E3",
            summary["year_salary"]
        ),
        (
            "G3",
            "Year Payable",
            "H3",
            summary["total_payable"]
        ),
    ]

    for label_cell, label, value_cell, value in info:
        ws[label_cell] = label
        ws[label_cell].font = Font(bold=True)
        ws[value_cell] = value
        ws[label_cell].border = border
        ws[value_cell].border = border

    # Monthly table
    headers = [
        "Month",
        "Days",
        "Recorded",
        "Present",
        "Absent",
        "CL",
        "SL",
        "Monthly Salary",
        "Absent Cut",
        "Payable"
    ]

    for col, value in enumerate(headers, 1):
        cell = ws.cell(5, col)
        cell.value = value
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center"
        )

    for row_no, item in enumerate(
        summary["monthly"],
        6
    ):
        values = [
            item["month"],
            item["days"],
            item["recorded"],
            item["present"],
            item["absent"],
            item["cl"],
            item["sl"],
            item["salary"],
            item["deduction"],
            item["payable"]
        ]

        for col, value in enumerate(values, 1):
            cell = ws.cell(
                row_no,
                col
            )
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center"
            )

            if col >= 8:
                cell.number_format = '₹#,##0.00'

    total_row = 18

    totals = [
        "YEAR TOTAL",
        "",
        summary["recorded"],
        summary["present"],
        summary["absent"],
        summary["cl"],
        summary["sl"],
        summary["year_salary"],
        summary["total_deduction"],
        summary["total_payable"]
    ]

    for col, value in enumerate(
        totals,
        1
    ):
        cell = ws.cell(
            total_row,
            col
        )
        cell.value = value
        cell.font = Font(bold=True)
        cell.fill = total_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center"
        )

        if col >= 8:
            cell.number_format = '₹#,##0.00'

    tab = Table(
        displayName="SalaryMonthlyTable",
        ref="A5:J17"
    )

    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    ws.add_table(tab)

    # Calculation note
    ws.merge_cells("A20:J20")
    ws["A20"] = "CALCULATION NOTE"
    ws["A20"].font = Font(
        bold=True,
        color="FFFFFF"
    )
    ws["A20"].fill = title_fill
    ws["A20"].alignment = Alignment(
        horizontal="center"
    )

    ws.merge_cells("A21:J21")
    ws["A21"] = (
        "Absent Cut = Absent Days × "
        "(Monthly Salary ÷ Calendar Days in Month). "
        "Present, CL and SL are treated as paid."
    )
    ws["A21"].alignment = Alignment(
        wrap_text=True
    )

    widths = [
        16, 10, 12, 12, 12,
        8, 8, 18, 16, 18
    ]

    for i, width in enumerate(
        widths,
        1
    ):
        ws.column_dimensions[
            get_column_letter(i)
        ].width = width

    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False

    # --------------------------------------------------------
    # Attendance Details
    # --------------------------------------------------------

    wd = wb.create_sheet(
        "Attendance Details"
    )

    detail_headers = [
        "Sr. No.",
        "Date",
        "Status"
    ]

    for col, value in enumerate(
        detail_headers,
        1
    ):
        cell = wd.cell(1, col)
        cell.value = value
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill
        cell.border = border

    attendance_rows = get_attendance(
        staff_id,
        year
    )

    for row_no, (
        date_value,
        status
    ) in enumerate(
        attendance_rows,
        2
    ):
        values = [
            row_no - 1,
            date_value,
            status
        ]

        for col, value in enumerate(
            values,
            1
        ):
            cell = wd.cell(
                row_no,
                col
            )
            cell.value = value
            cell.border = border

    for col, width in enumerate(
        [12, 18, 18],
        1
    ):
        wd.column_dimensions[
            get_column_letter(col)
        ].width = width

    wd.freeze_panes = "A2"
    wd.sheet_view.showGridLines = False

    # --------------------------------------------------------
    # Saved Salary Records
    # --------------------------------------------------------

    sr = wb.create_sheet(
        "Saved Salary Records"
    )

    saved_headers = [
        "Month",
        "Monthly Salary",
        "Calendar Days",
        "Present",
        "Absent",
        "CL",
        "SL",
        "Absent Cut",
        "Payable Salary"
    ]

    for col, value in enumerate(
        saved_headers,
        1
    ):
        cell = sr.cell(1, col)
        cell.value = value
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill
        cell.border = border

    saved = get_saved_records(
        staff_id,
        year
    )

    for row_no, row in enumerate(
        saved,
        2
    ):
        (
            month,
            salary,
            days,
            present,
            absent,
            cl,
            sl,
            cut,
            payable
        ) = row

        values = [
            calendar.month_name[month],
            salary,
            days,
            present,
            absent,
            cl,
            sl,
            cut,
            payable
        ]

        for col, value in enumerate(
            values,
            1
        ):
            cell = sr.cell(
                row_no,
                col
            )
            cell.value = value
            cell.border = border

            if col in (2, 8, 9):
                cell.number_format = '₹#,##0.00'

    for col, width in enumerate(
        [
            16, 18, 15, 12, 12,
            8, 8, 16, 18
        ],
        1
    ):
        sr.column_dimensions[
            get_column_letter(col)
        ].width = width

    sr.freeze_panes = "A2"
    sr.sheet_view.showGridLines = False

    wb.save(excel_path)


# ============================================================
# MAIN APPLICATION
# ============================================================

class SalaryReportApp:

    def __init__(self, root):
        self.root = root

        self.root.title(
            "Staff Management System - Salary Report"
        )

        self.root.geometry(
            "1400x900"
        )

        self.root.minsize(
            1150,
            720
        )

        self.root.configure(
            bg=BG
        )

        self.staff_var = tk.StringVar()
        self.year_var = tk.StringVar(
            value=str(datetime.now().year)
        )

        self.month_var = tk.StringVar(
            value=calendar.month_name[
                datetime.now().month
            ]
        )

        self.current_staff_id = None
        self.current_staff_name = ""
        self.current_salary = 0.0
        self.current_year = datetime.now().year
        self.current_month = datetime.now().month
        self.current_records = []
        self.current_summary = None
        self.current_month_item = None

        self.current_pdf = None
        self.current_excel = None

        setup_salary_records_table()

        self.build_ui()
        self.load_staff()

    # ========================================================
    # UI
    # ========================================================

    def build_ui(self):

        # Header
        header = tk.Frame(
            self.root,
            bg=HEADER,
            height=105
        )
        header.pack(
            fill="x"
        )
        header.pack_propagate(False)

        tk.Label(
            header,
            text="STAFF SALARY & PAYMENT REPORT",
            font=("Segoe UI", 24, "bold"),
            fg=WHITE,
            bg=HEADER
        ).pack(
            pady=(17, 2)
        )

        tk.Label(
            header,
            text=(
                "Monthly salary • Absent cut • "
                "Payable salary • Year total • "
                "Save record • Excel • Print"
            ),
            font=("Segoe UI", 10),
            fg="#bfdbfe",
            bg=HEADER
        ).pack()

        # Selection card
        select = tk.Frame(
            self.root,
            bg=WHITE,
            highlightthickness=1,
            highlightbackground=LINE
        )
        select.pack(
            fill="x",
            padx=22,
            pady=14
        )

        tk.Label(
            select,
            text="Staff",
            font=("Segoe UI", 10, "bold"),
            fg=TEXT,
            bg=WHITE
        ).grid(
            row=0,
            column=0,
            padx=(18, 8),
            pady=(12, 5),
            sticky="w"
        )

        tk.Label(
            select,
            text="Year",
            font=("Segoe UI", 10, "bold"),
            fg=TEXT,
            bg=WHITE
        ).grid(
            row=0,
            column=1,
            padx=8,
            pady=(12, 5),
            sticky="w"
        )

        tk.Label(
            select,
            text="Month",
            font=("Segoe UI", 10, "bold"),
            fg=TEXT,
            bg=WHITE
        ).grid(
            row=0,
            column=2,
            padx=8,
            pady=(12, 5),
            sticky="w"
        )

        self.staff_combo = ttk.Combobox(
            select,
            textvariable=self.staff_var,
            state="readonly",
            width=34,
            font=("Segoe UI", 10)
        )
        self.staff_combo.grid(
            row=1,
            column=0,
            padx=(18, 8),
            pady=(0, 14),
            ipady=5
        )

        self.year_combo = ttk.Combobox(
            select,
            textvariable=self.year_var,
            state="readonly",
            width=11,
            font=("Segoe UI", 10)
        )

        current_year = datetime.now().year

        years = [
            str(y)
            for y in range(
                current_year - 5,
                current_year + 6
            )
        ]

        self.year_combo["values"] = years

        self.year_combo.grid(
            row=1,
            column=1,
            padx=8,
            pady=(0, 14),
            ipady=5
        )

        months = [
            f"{i:02d} - {calendar.month_name[i]}"
            for i in range(1, 13)
        ]

        self.month_combo = ttk.Combobox(
            select,
            textvariable=self.month_var,
            state="readonly",
            width=20,
            font=("Segoe UI", 10)
        )

        self.month_combo["values"] = months

        self.month_combo.grid(
            row=1,
            column=2,
            padx=8,
            pady=(0, 14),
            ipady=5
        )

        # Buttons
        self.make_button(
            select,
            "GENERATE MONTH",
            BLUE,
            self.generate_month_report
        ).grid(
            row=1,
            column=3,
            padx=5,
            pady=(0, 14)
        )

        self.make_button(
            select,
            "GENERATE YEAR",
            PURPLE,
            self.generate_year_report
        ).grid(
            row=1,
            column=4,
            padx=5,
            pady=(0, 14)
        )

        self.make_button(
            select,
            "SAVE MONTH RECORD",
            GREEN,
            self.save_current_month
        ).grid(
            row=1,
            column=5,
            padx=5,
            pady=(0, 14)
        )

        self.make_button(
            select,
            "SAVE FULL YEAR",
            "#0f766e",
            self.save_full_year
        ).grid(
            row=1,
            column=6,
            padx=5,
            pady=(0, 14)
        )

        self.make_button(
            select,
            "EXCEL",
            "#15803d",
            self.export_excel
        ).grid(
            row=1,
            column=7,
            padx=5,
            pady=(0, 14)
        )

        self.make_button(
            select,
            "PRINT / PDF",
            ORANGE,
            self.print_report
        ).grid(
            row=1,
            column=8,
            padx=5,
            pady=(0, 14)
        )

        self.make_button(
            select,
            "REPORT FOLDER",
            "#e2e8f0",
            self.open_report_folder,
            TEXT
        ).grid(
            row=1,
            column=9,
            padx=(5, 18),
            pady=(0, 14)
        )

        # Report area
        outer = tk.Frame(
            self.root,
            bg=BG
        )
        outer.pack(
            fill="both",
            expand=True,
            padx=22,
            pady=(0, 18)
        )

        self.canvas = tk.Canvas(
            outer,
            bg=BG,
            highlightthickness=0
        )

        scrollbar = ttk.Scrollbar(
            outer,
            orient="vertical",
            command=self.canvas.yview
        )

        self.canvas.configure(
            yscrollcommand=scrollbar.set
        )

        scrollbar.pack(
            side="right",
            fill="y"
        )

        self.canvas.pack(
            side="left",
            fill="both",
            expand=True
        )

        self.report_frame = tk.Frame(
            self.canvas,
            bg=BG
        )

        self.canvas_window = self.canvas.create_window(
            (0, 0),
            window=self.report_frame,
            anchor="nw"
        )

        self.report_frame.bind(
            "<Configure>",
            lambda e: self.canvas.configure(
                scrollregion=self.canvas.bbox("all")
            )
        )

        self.canvas.bind(
            "<Configure>",
            lambda e: self.canvas.itemconfigure(
                self.canvas_window,
                width=e.width
            )
        )

        self.canvas.bind_all(
            "<MouseWheel>",
            self.mouse_scroll
        )

        tk.Label(
            self.report_frame,
            text=(
                "Select Staff + Year + Month and click "
                "GENERATE MONTH"
            ),
            font=("Segoe UI", 15, "bold"),
            fg=GRAY,
            bg=BG
        ).pack(
            pady=100
        )

    def mouse_scroll(self, event):
        self.canvas.yview_scroll(
            int(-1 * (event.delta / 120)),
            "units"
        )

    def make_button(
        self,
        parent,
        text,
        bg,
        command,
        fg=WHITE
    ):
        return tk.Button(
            parent,
            text=text,
            font=("Segoe UI", 9, "bold"),
            bg=bg,
            fg=fg,
            activebackground=bg,
            activeforeground=fg,
            relief="flat",
            cursor="hand2",
            padx=12,
            pady=9,
            command=command
        )

    # ========================================================
    # STAFF
    # ========================================================

    def load_staff(self):
        try:
            rows = get_staff_list()

            values = []

            for sid, name, salary in rows:
                try:
                    salary_value = float(
                        salary or 0
                    )
                except (ValueError, TypeError):
                    salary_value = 0

                values.append(
                    f"{sid} - {name} - "
                    f"Rs. {salary_value:,.2f}"
                )

            self.staff_combo["values"] = values

            if values:
                self.staff_combo.current(0)
            else:
                self.staff_combo.set(
                    "No Staff Available"
                )

        except Exception as e:
            messagebox.showerror(
                "Database Error",
                str(e),
                parent=self.root
            )

    def parse_selected_staff(self):
        selected = self.staff_var.get().strip()

        if " - " not in selected:
            raise ValueError(
                "Please select a staff member."
            )

        try:
            staff_id = int(
                selected.split(
                    " - ",
                    1
                )[0]
            )
        except ValueError:
            raise ValueError(
                "Invalid staff selection."
            )

        rows = get_staff_list()

        for sid, name, salary in rows:
            if sid == staff_id:
                try:
                    salary_value = float(
                        salary or 0
                    )
                except (ValueError, TypeError):
                    salary_value = 0

                return (
                    sid,
                    str(name),
                    salary_value
                )

        raise ValueError(
            "Selected staff was not found."
        )

    # ========================================================
    # MONTH / YEAR
    # ========================================================

    def get_selected_month(self):
        value = self.month_var.get().strip()

        if not value:
            return datetime.now().month

        try:
            return int(
                value.split(
                    " - ",
                    1
                )[0]
            )
        except (ValueError, IndexError):
            return datetime.now().month

    def load_selected_data(self):
        staff_id, staff_name, salary = (
            self.parse_selected_staff()
        )

        year = int(
            self.year_var.get()
        )

        month = self.get_selected_month()

        if salary <= 0:
            raise ValueError(
                "Selected staff has no valid monthly salary."
            )

        records = get_attendance(
            staff_id,
            year
        )

        summary = calculate_year(
            records,
            salary,
            year
        )

        month_item = calculate_month(
            records,
            salary,
            year,
            month
        )

        self.current_staff_id = staff_id
        self.current_staff_name = staff_name
        self.current_salary = salary
        self.current_year = year
        self.current_month = month
        self.current_records = records
        self.current_summary = summary
        self.current_month_item = month_item

    # ========================================================
    # GENERATE MONTH
    # ========================================================

    def generate_month_report(self):
        try:
            self.load_selected_data()
            self.show_month_report()

        except Exception as e:
            messagebox.showerror(
                "Report Error",
                str(e),
                parent=self.root
            )

    # ========================================================
    # GENERATE YEAR
    # ========================================================

    def generate_year_report(self):
        try:
            self.load_selected_data()
            self.show_year_report()

        except Exception as e:
            messagebox.showerror(
                "Report Error",
                str(e),
                parent=self.root
            )

    # ========================================================
    # CLEAR REPORT
    # ========================================================

    def clear_report(self):
        for widget in (
            self.report_frame.winfo_children()
        ):
            widget.destroy()

    # ========================================================
    # COMMON STAFF INFO
    # ========================================================

    def show_staff_info(self, title_text):
        info = tk.Frame(
            self.report_frame,
            bg=WHITE,
            highlightthickness=1,
            highlightbackground=LINE
        )

        info.pack(
            fill="x",
            pady=(0, 12)
        )

        tk.Label(
            info,
            text=title_text,
            font=("Segoe UI", 14, "bold"),
            fg=TEXT,
            bg=WHITE
        ).pack(
            side="left",
            padx=18,
            pady=14
        )

        tk.Label(
            info,
            text=(
                f"Staff: {self.current_staff_name}"
                f"   |   ID: {self.current_staff_id}"
                f"   |   Year: {self.current_year}"
            ),
            font=("Segoe UI", 10, "bold"),
            fg=TEXT,
            bg=WHITE
        ).pack(
            side="left",
            padx=18
        )

        tk.Label(
            info,
            text=(
                f"Monthly Salary: "
                f"Rs. {self.current_salary:,.2f}"
            ),
            font=("Segoe UI", 11, "bold"),
            fg=GREEN,
            bg=WHITE
        ).pack(
            side="right",
            padx=18
        )

    # ========================================================
    # MONTH REPORT
    # ========================================================

    def show_month_report(self):
        self.clear_report()

        item = self.current_month_item

        self.show_staff_info(
            f"{item['month'].upper()} SALARY REPORT"
        )

        # Summary cards
        cards = [
            (
                "MONTHLY SALARY",
                f"Rs. {item['salary']:,.2f}",
                BLUE
            ),
            (
                "PRESENT",
                item["present"],
                GREEN
            ),
            (
                "ABSENT",
                item["absent"],
                RED
            ),
            (
                "CL",
                item["cl"],
                BLUE
            ),
            (
                "SL",
                item["sl"],
                PURPLE
            ),
            (
                "ABSENT CUT",
                f"Rs. {item['deduction']:,.2f}",
                ORANGE
            ),
            (
                "PAYABLE",
                f"Rs. {item['payable']:,.2f}",
                GREEN
            ),
        ]

        card_frame = tk.Frame(
            self.report_frame,
            bg=BG
        )
        card_frame.pack(
            fill="x",
            pady=(0, 15)
        )

        for i, (
            title,
            value,
            color
        ) in enumerate(cards):

            card = tk.Frame(
                card_frame,
                bg=WHITE,
                highlightthickness=1,
                highlightbackground=LINE
            )

            card.grid(
                row=0,
                column=i,
                padx=4,
                sticky="nsew"
            )

            card_frame.columnconfigure(
                i,
                weight=1
            )

            tk.Label(
                card,
                text=title,
                font=("Segoe UI", 8, "bold"),
                fg=GRAY,
                bg=WHITE
            ).pack(
                pady=(10, 3)
            )

            tk.Label(
                card,
                text=str(value),
                font=("Segoe UI", 12, "bold"),
                fg=color,
                bg=WHITE
            ).pack(
                pady=(0, 10)
            )

        # Detailed monthly payment
        payment = tk.Frame(
            self.report_frame,
            bg=WHITE,
            highlightthickness=1,
            highlightbackground=LINE
        )
        payment.pack(
            fill="x",
            pady=(0, 15)
        )

        tk.Label(
            payment,
            text="MONTHLY PAYMENT CALCULATION",
            font=("Segoe UI", 15, "bold"),
            fg=TEXT,
            bg=WHITE
        ).pack(
            anchor="w",
            padx=20,
            pady=(15, 10)
        )

        rows = [
            (
                "Calendar Days",
                str(item["days"]),
                TEXT
            ),
            (
                "Attendance Records",
                str(item["recorded"]),
                TEXT
            ),
            (
                "Present Days",
                str(item["present"]),
                GREEN
            ),
            (
                "Absent Days",
                str(item["absent"]),
                RED
            ),
            (
                "CL Days",
                str(item["cl"]),
                BLUE
            ),
            (
                "SL Days",
                str(item["sl"]),
                PURPLE
            ),
            (
                "Daily Rate",
                f"Rs. {item['daily_rate']:,.2f}",
                TEXT
            ),
            (
                "Monthly Salary",
                f"Rs. {item['salary']:,.2f}",
                BLUE
            ),
            (
                "Absent Salary Cut",
                f"- Rs. {item['deduction']:,.2f}",
                RED
            ),
            (
                "FINAL PAYABLE SALARY",
                f"Rs. {item['payable']:,.2f}",
                GREEN
            ),
        ]

        for index, (
            label,
            value,
            color
        ) in enumerate(rows):

            bg = (
                LIGHT_GREEN
                if label == "FINAL PAYABLE SALARY"
                else WHITE
            )

            row = tk.Frame(
                payment,
                bg=bg
            )
            row.pack(
                fill="x",
                padx=20,
                pady=2
            )

            tk.Label(
                row,
                text=label,
                font=(
                    "Segoe UI",
                    10,
                    "bold"
                ),
                fg=TEXT,
                bg=bg,
                anchor="w"
            ).pack(
                side="left",
                padx=10,
                pady=7
            )

            tk.Label(
                row,
                text=value,
                font=(
                    "Segoe UI",
                    10,
                    "bold"
                ),
                fg=color,
                bg=bg
            ).pack(
                side="right",
                padx=10
            )

        # Formula
        formula = tk.Frame(
            self.report_frame,
            bg=LIGHT_BLUE
        )
        formula.pack(
            fill="x",
            pady=(0, 15)
        )

        tk.Label(
            formula,
            text=(
                f"Calculation: "
                f"Rs. {item['salary']:,.2f} ÷ "
                f"{item['days']} days = "
                f"Rs. {item['daily_rate']:,.2f} daily rate\n"
                f"Absent Cut: "
                f"{item['absent']} × "
                f"Rs. {item['daily_rate']:,.2f} = "
                f"Rs. {item['deduction']:,.2f}\n"
                f"Payable: "
                f"Rs. {item['salary']:,.2f} - "
                f"Rs. {item['deduction']:,.2f} = "
                f"Rs. {item['payable']:,.2f}"
            ),
            font=("Segoe UI", 10, "bold"),
            fg=TEXT,
            bg=LIGHT_BLUE,
            justify="left"
        ).pack(
            anchor="w",
            padx=18,
            pady=14
        )

        self.show_attendance_details(
            self.current_records,
            self.current_month
        )

        self.canvas.update_idletasks()
        self.canvas.configure(
            scrollregion=self.canvas.bbox("all")
        )

    # ========================================================
    # YEAR REPORT
    # ========================================================

    def show_year_report(self):
        self.clear_report()

        s = self.current_summary

        self.show_staff_info(
            f"{self.current_year} FULL YEAR SALARY REPORT"
        )

        # Year cards
        cards = [
            (
                "YEAR SALARY",
                f"Rs. {s['year_salary']:,.2f}",
                BLUE
            ),
            (
                "PRESENT",
                s["present"],
                GREEN
            ),
            (
                "ABSENT",
                s["absent"],
                RED
            ),
            (
                "CL",
                s["cl"],
                BLUE
            ),
            (
                "SL",
                s["sl"],
                PURPLE
            ),
            (
                "TOTAL CUT",
                f"Rs. {s['total_deduction']:,.2f}",
                ORANGE
            ),
            (
                "YEAR PAYABLE",
                f"Rs. {s['total_payable']:,.2f}",
                GREEN
            ),
        ]

        card_frame = tk.Frame(
            self.report_frame,
            bg=BG
        )
        card_frame.pack(
            fill="x",
            pady=(0, 15)
        )

        for i, (
            title,
            value,
            color
        ) in enumerate(cards):

            card = tk.Frame(
                card_frame,
                bg=WHITE,
                highlightthickness=1,
                highlightbackground=LINE
            )

            card.grid(
                row=0,
                column=i,
                padx=4,
                sticky="nsew"
            )

            card_frame.columnconfigure(
                i,
                weight=1
            )

            tk.Label(
                card,
                text=title,
                font=("Segoe UI", 8, "bold"),
                fg=GRAY,
                bg=WHITE
            ).pack(
                pady=(10, 3)
            )

            tk.Label(
                card,
                text=str(value),
                font=("Segoe UI", 12, "bold"),
                fg=color,
                bg=WHITE
            ).pack(
                pady=(0, 10)
            )

        # Year table
        tk.Label(
            self.report_frame,
            text="MONTH-WISE SALARY RECORD",
            font=("Segoe UI", 16, "bold"),
            fg=TEXT,
            bg=BG
        ).pack(
            anchor="w",
            pady=(0, 7)
        )

        table_card = tk.Frame(
            self.report_frame,
            bg=WHITE,
            highlightthickness=1,
            highlightbackground=LINE
        )
        table_card.pack(
            fill="x",
            pady=(0, 15)
        )

        headers = [
            "MONTH",
            "DAYS",
            "RECORDED",
            "PRESENT",
            "ABSENT",
            "CL",
            "SL",
            "SALARY",
            "ABSENT CUT",
            "PAYABLE"
        ]

        for c, h in enumerate(headers):
            tk.Label(
                table_card,
                text=h,
                font=("Segoe UI", 8, "bold"),
                bg=HEADER,
                fg=WHITE,
                borderwidth=1,
                relief="solid",
                padx=4,
                pady=8
            ).grid(
                row=0,
                column=c,
                sticky="nsew"
            )

        for r, item in enumerate(
            s["monthly"],
            1
        ):
            row_bg = (
                WHITE
                if r % 2
                else "#f8fafc"
            )

            values = [
                item["month"],
                item["days"],
                item["recorded"],
                item["present"],
                item["absent"],
                item["cl"],
                item["sl"],
                f"Rs. {item['salary']:,.0f}",
                f"Rs. {item['deduction']:,.0f}",
                f"Rs. {item['payable']:,.0f}"
            ]

            for c, value in enumerate(values):
                tk.Label(
                    table_card,
                    text=value,
                    font=(
                        "Segoe UI",
                        8,
                        "bold"
                        if c >= 3
                        else "normal"
                    ),
                    bg=row_bg,
                    fg=TEXT,
                    borderwidth=1,
                    relief="solid",
                    padx=4,
                    pady=7
                ).grid(
                    row=r,
                    column=c,
                    sticky="nsew"
                )

        for c in range(
            len(headers)
        ):
            table_card.columnconfigure(
                c,
                weight=1
            )

        # Year summary
        summary_card = tk.Frame(
            self.report_frame,
            bg=WHITE,
            highlightthickness=1,
            highlightbackground=LINE
        )

        summary_card.pack(
            fill="x",
            pady=(0, 15)
        )

        tk.Label(
            summary_card,
            text="YEARLY PAYMENT SUMMARY",
            font=("Segoe UI", 15, "bold"),
            fg=TEXT,
            bg=WHITE
        ).pack(
            anchor="w",
            padx=20,
            pady=(15, 10)
        )

        summary_text = (
            f"Annual Salary: "
            f"Rs. {s['year_salary']:,.2f}\n"
            f"Present: {s['present']}    |    "
            f"Absent: {s['absent']}    |    "
            f"CL: {s['cl']}    |    "
            f"SL: {s['sl']}    |    "
            f"Recorded: {s['recorded']}\n"
            f"Total Absent Salary Cut: "
            f"Rs. {s['total_deduction']:,.2f}\n"
            f"FINAL YEAR PAYMENT: "
            f"Rs. {s['total_payable']:,.2f}"
        )

        tk.Label(
            summary_card,
            text=summary_text,
            font=("Segoe UI", 11, "bold"),
            fg=TEXT,
            bg=WHITE,
            justify="left"
        ).pack(
            anchor="w",
            padx=20,
            pady=(0, 18)
        )

        tk.Label(
            self.report_frame,
            text=(
                "Calculation: Every month starts with the full monthly "
                "salary. Only Absent days are deducted. "
                "Present, CL and SL are paid."
            ),
            font=("Segoe UI", 9),
            fg=GRAY,
            bg=BG
        ).pack(
            anchor="w",
            pady=(0, 20)
        )

        self.canvas.update_idletasks()
        self.canvas.configure(
            scrollregion=self.canvas.bbox("all")
        )

    # ========================================================
    # ATTENDANCE DETAILS
    # ========================================================

    def show_attendance_details(
        self,
        records,
        month
    ):
        title = tk.Label(
            self.report_frame,
            text=(
                f"{calendar.month_name[month].upper()} "
                f"ATTENDANCE DETAILS"
            ),
            font=("Segoe UI", 15, "bold"),
            fg=TEXT,
            bg=BG
        )

        title.pack(
            anchor="w",
            pady=(0, 7)
        )

        card = tk.Frame(
            self.report_frame,
            bg=WHITE,
            highlightthickness=1,
            highlightbackground=LINE
        )

        card.pack(
            fill="x",
            pady=(0, 20)
        )

        headers = [
            "SR. NO.",
            "DATE",
            "STATUS"
        ]

        for c, h in enumerate(headers):
            tk.Label(
                card,
                text=h,
                font=("Segoe UI", 9, "bold"),
                bg=HEADER,
                fg=WHITE,
                borderwidth=1,
                relief="solid",
                padx=10,
                pady=8
            ).grid(
                row=0,
                column=c,
                sticky="nsew"
            )

        month_records = []

        for date_value, status in records:
            try:
                dt = datetime.strptime(
                    str(date_value),
                    "%Y-%m-%d"
                )

                if dt.month == month:
                    month_records.append(
                        (
                            date_value,
                            status
                        )
                    )
            except ValueError:
                pass

        if not month_records:
            tk.Label(
                card,
                text="No attendance record found for this month.",
                font=("Segoe UI", 10),
                fg=GRAY,
                bg=WHITE
            ).grid(
                row=1,
                column=0,
                columnspan=3,
                pady=18
            )
            return

        for r, (
            date_value,
            status
        ) in enumerate(
            month_records,
            1
        ):
            bg = (
                WHITE
                if r % 2
                else "#f8fafc"
            )

            values = [
                r,
                date_value,
                status
            ]

            for c, value in enumerate(values):
                tk.Label(
                    card,
                    text=value,
                    font=("Segoe UI", 9, "bold"),
                    bg=bg,
                    fg=TEXT,
                    borderwidth=1,
                    relief="solid",
                    padx=10,
                    pady=7
                ).grid(
                    row=r,
                    column=c,
                    sticky="nsew"
                )

        for c in range(3):
            card.columnconfigure(
                c,
                weight=1
            )

    # ========================================================
    # SAVE MONTH RECORD
    # ========================================================

    def save_current_month(self):
        try:
            self.load_selected_data()

            if not self.current_month_item:
                raise ValueError(
                    "Generate the month report first."
                )

            save_month_record(
                self.current_staff_id,
                self.current_year,
                self.current_month_item
            )

            messagebox.showinfo(
                "Record Saved",
                (
                    "Monthly salary record saved successfully.\n\n"
                    f"Staff: {self.current_staff_name}\n"
                    f"Month: "
                    f"{self.current_month_item['month']} "
                    f"{self.current_year}\n"
                    f"Monthly Salary: "
                    f"Rs. {self.current_month_item['salary']:,.2f}\n"
                    f"Absent Cut: "
                    f"Rs. {self.current_month_item['deduction']:,.2f}\n"
                    f"Payable: "
                    f"Rs. {self.current_month_item['payable']:,.2f}"
                ),
                parent=self.root
            )

        except Exception as e:
            messagebox.showerror(
                "Save Error",
                str(e),
                parent=self.root
            )

    # ========================================================
    # SAVE FULL YEAR
    # ========================================================

    def save_full_year(self):
        try:
            self.load_selected_data()

            save_year_records(
                self.current_staff_id,
                self.current_year,
                self.current_summary
            )

            messagebox.showinfo(
                "Year Records Saved",
                (
                    "All 12 month salary records saved successfully.\n\n"
                    f"Staff: {self.current_staff_name}\n"
                    f"Year: {self.current_year}\n"
                    f"Annual Salary: "
                    f"Rs. {self.current_summary['year_salary']:,.2f}\n"
                    f"Total Absent Cut: "
                    f"Rs. "
                    f"{self.current_summary['total_deduction']:,.2f}\n"
                    f"Final Year Payment: "
                    f"Rs. "
                    f"{self.current_summary['total_payable']:,.2f}"
                ),
                parent=self.root
            )

        except Exception as e:
            messagebox.showerror(
                "Save Year Error",
                str(e),
                parent=self.root
            )

    # ========================================================
    # FILE PATHS
    # ========================================================

    def ensure_report_files(self):
        if not self.current_summary:
            self.load_selected_data()

        if not self.current_summary:
            return False

        safe_name = "".join(
            ch
            if ch.isalnum() or ch in " _-"
            else "_"
            for ch in self.current_staff_name
        ).strip() or "Staff"

        staff_dir = os.path.join(
            REPORT_DIR,
            safe_name
        )

        os.makedirs(
            staff_dir,
            exist_ok=True
        )

        self.current_pdf = os.path.join(
            staff_dir,
            f"Salary_Report_"
            f"{self.current_staff_id}_"
            f"{self.current_year}.pdf"
        )

        self.current_excel = os.path.join(
            staff_dir,
            f"Salary_Report_"
            f"{self.current_staff_id}_"
            f"{self.current_year}.xlsx"
        )

        return True

    # ========================================================
    # EXCEL
    # ========================================================

    def export_excel(self):
        try:
            self.load_selected_data()

            if not self.ensure_report_files():
                return

            create_excel(
                self.current_staff_name,
                self.current_staff_id,
                self.current_year,
                self.current_salary,
                self.current_summary,
                self.current_excel
            )

            messagebox.showinfo(
                "Excel Generated",
                (
                    "Professional salary spreadsheet generated "
                    "successfully.\n\n"
                    f"{self.current_excel}"
                ),
                parent=self.root
            )

            self.open_file(
                self.current_excel
            )

        except Exception as e:
            messagebox.showerror(
                "Excel Error",
                str(e),
                parent=self.root
            )

    # ========================================================
    # PRINT / PDF
    # ========================================================

    def print_report(self):
        try:
            self.load_selected_data()

            if not self.ensure_report_files():
                return

            create_pdf(
                self.current_staff_name,
                self.current_staff_id,
                self.current_year,
                self.current_salary,
                self.current_summary,
                self.current_pdf
            )

            messagebox.showinfo(
                "PDF Generated",
                (
                    "Salary PDF generated successfully.\n\n"
                    "The PDF will open now. "
                    "Press Ctrl + P to print it."
                ),
                parent=self.root
            )

            self.open_file(
                self.current_pdf
            )

        except Exception as e:
            messagebox.showerror(
                "PDF Error",
                str(e),
                parent=self.root
            )

    # ========================================================
    # OPEN FILE
    # ========================================================

    def open_file(self, path):
        if not os.path.exists(path):
            raise FileNotFoundError(
                path
            )

        if os.name == "nt":
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(
                ["open", path]
            )
        else:
            subprocess.Popen(
                ["xdg-open", path]
            )

    # ========================================================
    # OPEN REPORT FOLDER
    # ========================================================

    def open_report_folder(self):
        try:
            if os.name == "nt":
                os.startfile(
                    REPORT_DIR
                )
            elif sys.platform == "darwin":
                subprocess.Popen(
                    ["open", REPORT_DIR]
                )
            else:
                subprocess.Popen(
                    ["xdg-open", REPORT_DIR]
                )

        except Exception as e:
            messagebox.showerror(
                "Folder Error",
                str(e),
                parent=self.root
            )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    try:
        setup_salary_records_table()

        root = tk.Tk()

        app = SalaryReportApp(
            root
        )

        root.mainloop()

    except Exception as e:
        try:
            messagebox.showerror(
                "Application Error",
                str(e)
            )
        except Exception:
            print(e)
